"""Inspect authoritative data and real exports; never infer browser success."""
import csv
import hashlib
import json
import subprocess
import shutil
import xml.etree.ElementTree as ET
from pathlib import Path
from PIL import Image,ImageDraw
from pptx import Presentation
import numpy as np
from workshop.experiments import FUNCTIONS, ALGORITHMS, PARAMETERS, save_json, summaries, source_hashes


def check(condition,message):
    if not condition:raise AssertionError(message)


def render_pdf(pdf,folder):
    folder.mkdir(parents=True,exist_ok=True)
    subprocess.run(['pdftoppm','-png','-scale-to','1600',str(pdf),str(folder/'slide')],check=True,capture_output=True)
    paths=sorted(folder.glob('slide-*.png'));check(len(paths)==12,'PDF must render exactly 12 pages')
    sheet=Image.new('RGB',(1280,6*385),'#dce2e8');draw=ImageDraw.Draw(sheet)
    for i,p in enumerate(paths):
        with Image.open(p) as original:
            original.verify()
        with Image.open(p) as original:
            im=original.copy();im.thumbnail((620,350));x=10+(i%2)*640;y=27+(i//2)*385
            sheet.paste(im,(x,y));draw.text((x,y-19),f'Slide {i+1:02d}',fill='black')
    sheet.save(folder/'contact_sheet.png')
    return paths


def verify_deliverables(root=Path('results/workshop'),presentation=Path('presentation')):
    root=Path(root);presentation=Path(presentation)
    config=json.loads((root/'config.json').read_text());records=json.loads((root/'runs.json').read_text())
    rows=json.loads((root/'summary.json').read_text());provenance=json.loads((root/'provenance.json').read_text())
    expected_runs=config['runs_per_algorithm_function']*6
    check(len(records)==expected_runs,'Wrong statistical run count')
    check(provenance['source_sha256']==source_hashes(),'Experiment source changed after raw data generation: regenerate experiment')
    seen=set();raw_fitness_values=0
    for r in records:
        key=(r['function'],r['algorithm'],r['run']);check(key not in seen,'Duplicate run');seen.add(key)
        fi=list(FUNCTIONS).index(r['function']);expected_seed=config['base_seed']+10000*fi+r['run']-1
        check(r['seed']==expected_seed,'Seed schedule mismatch')
        check(r['parameters']==PARAMETERS[r['algorithm']],'Parameters mismatch')
        func,bounds,optimum=FUNCTIONS[r['function']]
        check(r['lower_bound']==[bounds[0]]*2 and r['upper_bound']==[bounds[1]]*2,'Bounds mismatch')
        check(r['dimensions']==2 and r['agents']==config['agents'] and r['iterations']==config['iterations'],'Run config mismatch')
        check(func(optimum)==0.,'Incorrect known optimum')
        with np.load(root/r['trace_file']) as trace:
            positions=trace['positions'];fitness=trace['population_fitness'];best=trace['best_fitness'];best_positions=trace['best_positions'];calls=trace['evaluations']
        check(positions.shape==(config['iterations']+1,config['agents'],2),'Frame dimensions mismatch')
        check(fitness.shape==positions.shape[:2],'Population fitness not aligned')
        check(len(best)==len(best_positions)==len(calls)==len(positions),'Archive/history alignment mismatch')
        check(np.all((positions>=bounds[0])&(positions<=bounds[1])),'Agents outside bounds')
        # Independent re-evaluation of every saved population point and archive.
        if r['function']=='sphere':
            actual=(positions**2).sum(axis=2);actual_best=(best_positions**2).sum(axis=1)
        else:
            x,y=positions[:,:,0],positions[:,:,1];actual=100*(y-x*x)**2+(x-1)**2
            x,y=best_positions[:,0],best_positions[:,1];actual_best=100*(y-x*x)**2+(x-1)**2
        np.testing.assert_allclose(actual,fitness,rtol=1e-12,atol=1e-15)
        np.testing.assert_allclose(actual_best,best,rtol=1e-12,atol=1e-15)
        np.testing.assert_allclose(best,np.minimum.accumulate(fitness.min(axis=1)),rtol=1e-12,atol=0)
        np.testing.assert_array_equal(best,r['convergence_history'])
        np.testing.assert_array_equal(best_positions,r['best_positions_history'])
        np.testing.assert_array_equal(calls,config['agents']*np.arange(1,config['iterations']+2))
        check(r['evaluation_count']==int(calls[-1]) and r['final_best_fitness']==float(best[-1]),'Wrong final result')
        check(r['best_position']==best_positions[-1].tolist(),'Wrong final position')
        check(r['runtime_seconds']>0,'Invalid runtime')
        single=json.loads((root/'runs'/f"{r['function']}_{r['algorithm']}_{r['run']:02d}.json").read_text())
        check(single==r,'Aggregate and per-run JSON differ')
        with (root/r['history_csv']).open() as handle: csvrows=list(csv.DictReader(handle))
        check(len(csvrows)==len(best),'CSV initial frame missing')
        for t,row in enumerate(csvrows):
            check(int(row['iteration'])==t and float(row['best_so_far'])==best[t] and int(row['evaluations'])==calls[t],'CSV trace mismatch')
        raw_fitness_values+=fitness.size
    recalculated=summaries(records)
    check(recalculated==rows,'Summary is not computed from all current raw runs')
    with (root/'summary.csv').open() as handle: csvsummary=list(csv.DictReader(handle))
    for a,b in zip(rows,csvsummary):
        for k,v in a.items():check(b[k]==str(v),'Summary CSV mismatch')
    with (root/'runs.csv').open() as handle: csvruns=list(csv.DictReader(handle))
    check(len(csvruns)==len(records),'Global run CSV wrong count')
    for a,b in zip(records,csvruns):
        for k,v in a.items():
            if isinstance(v,(list,dict)):check(json.loads(b[k])==v,'Global CSV nested result mismatch')
            else:check(b[k]==str(v),'Global CSV result mismatch')
    # Same initial states across methods, distinct streams across runs.
    for name in FUNCTIONS:
        for run in range(1,config['runs_per_algorithm_function']+1):
            initial=[np.load(root/'runs'/f'{name}_{a}_{run:02d}.npz')['positions'][0] for a in ALGORITHMS]
            for p in initial[1:]:np.testing.assert_array_equal(p,initial[0])
    images_checked=0;svgs_checked=0
    for folder in ['plots','snapshots']:
        files=list((root/folder).glob('*.png'));check(bool(files),'Missing raster charts')
        for p in files:
            with Image.open(p) as im:
                check(im.width>=1500,'PNG export too small');im.verify()
            ET.parse(p.with_suffix('.svg'));svgs_checked+=1;images_checked+=1
    for name in FUNCTIONS:
        for kind in ['mean','median','gap','evaluations','population_mean','boxplot','table']:
            check((root/'plots'/f'{name}_{kind}.png').exists(),'Missing required statistical plot')
        for t in [0,25,50,75,100]:check((root/'snapshots'/f'{name}_iter{t:03d}.png').exists(),'Missing requested snapshot')
        example=json.loads((root/'examples'/f'{name}_goa_00.json').read_text())
        check(example['seed']==config['animation_seed']==424242,'Illustration seed changed')
        # Recompute the preselected run and compare every frame, not only endpoint.
        from algorithms.core import optimize
        function,bounds,_=FUNCTIONS[name]
        rerun=optimize('goa',function,2,[bounds[0]]*2,[bounds[1]]*2,config['agents'],config['iterations'],424242,**PARAMETERS['goa'])
        with np.load(root/example['trace_file']) as trace:
            np.testing.assert_array_equal(trace['positions'],rerun.positions_history)
            np.testing.assert_array_equal(trace['best_fitness'],rerun.history)
            np.testing.assert_array_equal(trace['best_positions'],rerun.best_positions_history)
    videos=[]
    for name in FUNCTIONS:
        gif=root/'animations'/f'{name}_goa.gif'
        with Image.open(gif) as im:
            check(im.n_frames==101,'GIF frame count mismatch');im.seek(100);im.load()
            gif_size=list(im.size)
        mp4=root/'animations'/f'{name}_goa.mp4'
        if mp4.exists():
            probe=json.loads(subprocess.check_output(['ffprobe','-v','error','-select_streams','v:0','-show_entries','stream=width,height,nb_frames,duration,codec_name','-of','json',str(mp4)]))['streams'][0]
            check(int(probe['nb_frames'])==101,'MP4 frame count mismatch')
            subprocess.run(['ffmpeg','-v','error','-i',str(mp4),'-f','null','-'],check=True,capture_output=True)
        else:probe={'fallback':'GIF only','reason':json.loads((root/'animations'/'manifest.json').read_text())}
        videos.append(dict(function=name,gif_frames=101,gif_size=gif_size,mp4=probe))
    deck=Presentation(presentation/'GOA_Optimization_Workshop.pptx')
    check(len(deck.slides)==12,'Wrong PPTX slide count')
    for i,slide in enumerate(deck.slides):
        for shape in slide.shapes:
            check(shape.left>=0 and shape.top>=0,'Slide shape outside top/left')
            check(shape.left+shape.width<=deck.slide_width+1000 and shape.top+shape.height<=deck.slide_height+1000,'Slide shape outside page')
        check('Allocated time:' in slide.notes_slide.notes_text_frame.text,'Missing speaker timing notes')
    plan=json.loads((presentation/'slide_plan.json').read_text());check(plan['total_seconds']==sum(s['seconds'] for s in plan['slides'])==720,'Wrong talk timing')
    check((presentation/'script_th.md').exists() and (presentation/'qa_th.md').read_text().count('\n## ')>=8,'Missing script/Q&A')
    pdf=presentation/'GOA_Optimization_Workshop.pdf';check(pdf.exists(),'Missing PDF')
    conversion=json.loads((presentation/'export_status.json').read_text())
    check(conversion['pdf_available'] and conversion.get('returncode')==0,'PDF export did not succeed for the current deck')
    check(conversion['pptx_sha256']==hashlib.sha256((presentation/'GOA_Optimization_Workshop.pptx').read_bytes()).hexdigest(),'PPTX changed after PDF export')
    check(conversion['pdf_sha256']==hashlib.sha256(pdf.read_bytes()).hexdigest(),'PDF changed after export')
    info=subprocess.check_output(['pdfinfo',str(pdf)],text=True);check('Pages:           12' in info,'Wrong PDF page count')
    text=subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True)
    check(text.count('\f')==12,'PDF page text missing')
    for a in rows:
        check(f"{a['runtime_mean_s']:.4f}" in text,'Measured runtime missing in PDF')
        for k in ['mean','median','sd','best','worst']:check(f"{a[k]:.2e}" in text,'Measured statistic missing/clipped in PDF')
    for term in ['Grasshopper','Eq. 2.7/2.8','3,030','3030','Rosenbrock','Sample','sample','10.1016/j.advengsoft.2017.01.004']:
        if term=='3030':continue
        if term=='Sample':continue
        check(term in text,f'Expected PDF content missing: {term}')
    (presentation/'rendered').mkdir(exist_ok=True)
    rendered=render_pdf(pdf,presentation/'rendered')
    # Word boxes from the actual PDF are checked against the slide dimensions.
    bbox=subprocess.check_output(['pdftotext','-bbox',str(pdf),'-'],text=True)
    doc=ET.fromstring(bbox);outside=[]
    for page in doc.iter():
        if not page.tag.endswith('page'):continue
        width,height=float(page.attrib['width']),float(page.attrib['height'])
        for word in page.iter():
            if word.tag.endswith('word'):
                if float(word.attrib['xMin'])<0 or float(word.attrib['yMin'])<0 or float(word.attrib['xMax'])>width or float(word.attrib['yMax'])>height:
                    outside.append(word.text)
    check(not outside,f'PDF text outside page: {outside}')
    report=dict(status='passed',statistical_runs=len(records),recalculated_population_fitness_values=raw_fitness_values,
                trace_states_per_run=config['iterations']+1,agents_per_state=config['agents'],
                all_raw_frames_and_best_positions_recomputed=True,all_run_config_and_seeds_verified=True,
                all_summaries_recomputed_from_every_run=True,evaluations_per_run=config['agents']*(config['iterations']+1),
                example_traces_reproduced=True,png_exports_checked=images_checked,svg_exports_checked=svgs_checked,
                videos=videos,pptx_slides=12,pdf_pages=12,pdf_text_inside_page=True,
                exact_measured_statistics_found_in_pdf=True,talk_allocated_seconds=720,
                source_hashes_current=source_hashes(),browser_verification='See docs/verification/browser.json; separate real-browser test',
                visual_review='See rendered slide pages and contact sheet; automated geometry alone does not prove readability')
    out=Path('docs/verification');out.mkdir(parents=True,exist_ok=True);save_json(out/'deliverables.json',report)
    print(json.dumps(report,indent=2))
    return report
