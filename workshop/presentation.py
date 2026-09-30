"""Generate a 12-slide, 12-minute deck directly from measured results."""
import json
import shutil
import subprocess
import tempfile
import hashlib
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from workshop.experiments import save_json

NAVY='152B46'; TEAL='087E8B'; INK='20354B'; MUTED='536477'; BG='F5F7FA'; GOLD='D76A18'
FONT='IBM Plex Sans Thai'
TIMES=[60,75,65,50,70,75,90,70,70,50,35,10]
TITLES=['GOA: from swarm to measured evidence','The GOA update rule','Exploration → exploitation','Where GOA can be applied','A reproducible comparison','Sphere: a smooth bowl','Rosenbrock: a narrow curved valley','Mean convergence and optimality gap','Typical performance and run variability','What the 180 runs show','Limits and conclusions','References & reproducibility']


def rgb(color): return RGBColor.from_string(color)


def textbox(slide,text,x,y,w,h,size=24,color=INK,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02)
    tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT
        p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(9)
    return shape


def rect(slide,x,y,w,h,color):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid();s.fill.fore_color.rgb=rgb(color);s.line.fill.background();return s


def picture_fit(slide,path,x,y,w,h):
    from PIL import Image
    with Image.open(path) as im: iw,ih=im.size
    factor=min(w/iw,h/ih);dw,dh=iw*factor,ih*factor
    return slide.shapes.add_picture(str(path),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),Inches(dw),Inches(dh))


def build_presentation(root=Path('results/workshop'),output=Path('presentation')):
    root=Path(root);output=Path(output);output.mkdir(parents=True,exist_ok=True)
    rows=json.loads((root/'summary.json').read_text());config=json.loads((root/'config.json').read_text())
    if config['runs_per_algorithm_function']!=30 or config['iterations']!=100 or config['agents']!=30:
        raise ValueError('Workshop deck requires the specified 30 runs, 30 agents, 100 iterations')
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    prs.core_properties.title='GOA — Optimization Workshop / Aj. Eckart'
    prs.core_properties.subject='Reproducible Sphere and Rosenbrock comparison with GA and RLS'
    prs.core_properties.author='Optimization Workshop'
    for i,title in enumerate(TITLES):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
        rect(s,0,0,13.333,.16,TEAL);textbox(s,title,.55,.38,12.1,.8,34,NAVY,True)
        textbox(s,f'Optimization Workshop · Aj. Eckart  |  {i+1:02d}/12',.6,7.09,10,.22,11,MUTED)
        textbox(s,f'{TIMES[i]} s',12,7.07,.7,.25,11,MUTED)
    s=prs.slides[0]
    textbox(s,'Grasshopper Optimisation Algorithm',.7,1.55,11.6,.8,36,TEAL,True)
    textbox(s,'ทฤษฎี • ภาพการเคลื่อนที่ • ผลทดลองที่ทำซ้ำได้',.72,2.5,11.7,.55,26)
    textbox(s,'Can social interactions find a good minimum?\nCompare GOA with real-number GA and parallel RLS.\nTwo functions · 180 independent measured runs',.72,3.55,11.8,2.25,27)
    rect(s,.72,6.1,11.85,.55,NAVY);textbox(s,'12-minute talk  |  2D visualization  |  All raw runs retained',1,6.17,11.2,.4,20,'FFFFFF')
    s=prs.slides[1]
    textbox(s,'Social interaction',.72,1.43,5.5,.5,25,TEAL,True)
    textbox(s,'s(r) = 0.5 exp(−r/1.5) − exp(−r)',.72,2.06,11.8,.6,29)
    textbox(s,'xᵢᵈ(t) = Tᵈ(t−1) + cₜ Σⱼ≠ᵢ [ cₜ (ubᵈ−lbᵈ)/2 · s(rᵢⱼᵈ) · uᵢⱼᵈ ]',.72,3.0,11.85,1.12,25,NAVY,True)
    textbox(s,'cₜ = 1 − t (1 − 10⁻⁵) / 100;   t = 1, …, 100',.72,4.35,11.8,.65,27)
    textbox(s,'T = best-so-far archive; u = unit direction from i to j.\nEq. 2.7/2.8 of Saremi et al. (2017), with a documented normalization choice.',.72,5.32,11.8,1.15,21,MUTED)
    s=prs.slides[2]
    for x,title,body,col in [(0.72,'Early: explore','Higher c → stronger interactions\nRepulsion spreads nearby agents\nAttraction links distant agents',TEAL),(6.85,'Late: exploit','Lower c → smaller movement\nSearch around the best archive\nClustering can still miss the optimum',GOLD)]:
        rect(s,x,1.6,5.75,3.15,'FFFFFF');textbox(s,title,x+.25,1.85,5.25,.65,28,col,True);textbox(s,body,x+.25,2.75,5.25,1.8,24)
    textbox(s,'rᵢⱼᵈ = 1 + 3 |xⱼᵈ − xᵢᵈ| / (ubᵈ − lbᵈ) ∈ [1, 4]',.72,5.08,11.8,.65,27)
    textbox(s,'This coordinate/bound mapping is our explicit implementation choice.\nSynchronous updates · clip to bounds · both c factors retained · no claim of exact paper replication',.72,5.95,11.8,.85,19,MUTED)
    s=prs.slides[3]
    textbox(s,'Structural design: minimize weight subject to stress limits',.72,1.55,11.8,1,28,TEAL,True)
    textbox(s,'Design variables → cross-sectional areas\nObjective → material / structural weight\nConstraints → stress or deflection limits',.72,2.85,11.8,2.1,26)
    textbox(s,'The original paper studies truss and cantilever design.\nOur workshop tests are unconstrained mathematical functions;\nwe do not claim new engineering results.',.72,5.25,11.8,1.25,23,MUTED)
    s=prs.slides[4]
    textbox(s,'Sphere: [−5, 5]²          Rosenbrock: [−2, 2]²',.72,1.55,11.85,.7,28,TEAL,True)
    textbox(s,'2 dimensions · 30 agents · 100 iterations · 30 runs per method\nSeeds: 20260930 + 10000 × function index + run index\nSame initial points for all three methods within a function/run\n3,030 actual objective calls per run: 30 + 100 × 30',.72,2.55,11.85,2.65,24)
    textbox(s,'GA: pc = 0.8, pm = 0.1; mutation σ = 0.1 × bound width; no elitism\nRLS: Gaussian step σ = 0.3; accept only non-worse moves\nAll methods archive best-so-far; time excludes graphics/export',.72,5.5,11.85,1.15,20,MUTED)
    for i,name in [(5,'sphere'),(6,'rosenbrock')]:
        s=prs.slides[i]
        formula='f(x) = x₁² + x₂²; optimum (0, 0), f* = 0' if name=='sphere' else 'f(x) = 100(x₂ − x₁²)² + (x₁ − 1)²; optimum (1, 1), f* = 0'
        textbox(s,formula,.72,1.38,11.8,.6,25,TEAL,True)
        picture_fit(s,root/'snapshots'/f'{name}_slide.png',.62,2.05,12.05,3.85)
        record=json.loads((root/'examples'/f'{name}_goa_00.json').read_text())
        textbox(s,f"Single illustration · preselected seed 424242 · final best = {record['final_best_fitness']:.3e}\nGreen star = known optimum; orange diamond = archive; red points = current agents",.72,6.03,11.9,.8,18,MUTED)
        # Slide hyperlinks to portable media; snapshots are PDF fallback.
        link=textbox(s,'▶ Open GIF / MP4 (16.8 s)',9.2,6.77,3.25,.25,14,GOLD,True)
        run=link.text_frame.paragraphs[0].runs[0]
        run.hyperlink.address=f'../results/workshop/animations/{name}_goa.mp4'
        if name=='rosenbrock':textbox(s,'Contour: log10(1 + f)',.72,5.8,8,.3,16,MUTED)
    s=prs.slides[7]
    picture_fit(s,root/'plots'/'sphere_mean_slide.png',.55,1.45,6.05,4.75)
    picture_fit(s,root/'plots'/'rosenbrock_gap_slide.png',6.7,1.45,6.05,4.75)
    textbox(s,'Each line averages best-so-far over all 30 independent runs.\nShading = ± sample SD. Gap = fitness here (f* = 0); log display floor = 10⁻¹⁶.',.72,6.12,11.8,.7,20,MUTED)
    s=prs.slides[8]
    picture_fit(s,root/'plots'/'sphere_median_slide.png',.55,1.45,6.05,4.75)
    picture_fit(s,root/'plots'/'rosenbrock_boxplot_slide.png',6.7,1.45,6.05,4.75)
    textbox(s,'GOA has the lowest median on both functions in this setup.\nRosenbrock: GOA’s long upper tail means some runs still stall.',.72,6.12,11.8,.7,21,MUTED)
    s=prs.slides[9]
    textbox(s,'Final fitness ↓ · all 30 runs per method · lower is better',.72,1.45,11.8,.5,23,TEAL,True)
    headers=['Function / method','Mean ± SD','Median','Best','Worst','Time (s)']
    table=s.shapes.add_table(7,6,Inches(.72),Inches(2.15),Inches(11.9),Inches(3.9)).table
    widths=[2.5,2.65,1.75,1.65,1.65,1.7]
    for i,w in enumerate(widths):table.columns[i].width=Inches(w)
    for j,value in enumerate(headers):table.cell(0,j).text=value
    for i,r in enumerate(rows,1):
        values=[f"{r['function'].title()} / {r['algorithm'].upper()}",f"{r['mean']:.2e} ± {r['sd']:.2e}",f"{r['median']:.2e}",f"{r['best']:.2e}",f"{r['worst']:.2e}",f"{r['runtime_mean_s']:.4f}"]
        for j,value in enumerate(values):table.cell(i,j).text=value
    for i in range(7):
        for j in range(6):
            cell=table.cell(i,j);cell.fill.solid();cell.fill.fore_color.rgb=rgb(NAVY if i==0 else ('FFFFFF' if i%2 else 'E9EFF4'))
            cell.margin_left=cell.margin_right=Inches(.08);cell.margin_top=Inches(.05)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(17 if j==1 else 18);p.font.color.rgb=rgb('FFFFFF' if i==0 else INK);p.font.bold=(i==0)
    textbox(s,'Equal measured budgets: 3,030 calls/run. Time = mean search time on this machine.\nRLS is faster here; GOA has the lowest median, but Rosenbrock worst-case GOA > RLS.',.72,6.22,11.8,.65,18,MUTED)
    s=prs.slides[10]
    textbox(s,'GOA performs well here; there is no universal winner.',.72,1.55,11.8,.75,29,TEAL,True)
    textbox(s,'Only two smooth 2D functions; no real design constraints tested\nFixed parameters; no tuning study or significance claim\nGOA interactions: O(T N² D); may stagnate near an imperfect target\nSwarm clustering ≠ proof that the known optimum was found',.72,2.75,11.8,2.55,25)
    textbox(s,'Next: higher dimensions, multimodal functions, sensitivity studies,\nand constrained applications with explicit feasibility checks.',.72,5.7,11.8,1.1,23,MUTED)
    s=prs.slides[11]
    textbox(s,'Saremi, Mirjalili & Lewis (2017). Grasshopper Optimisation Algorithm:\nTheory and application. Advances in Engineering Software, 105, 30–47.\nDOI: 10.1016/j.advengsoft.2017.01.004',.72,1.5,11.85,1.65,23)
    textbox(s,'Course material supplied in repository: workshop_objective_function.pdf\nGA lecture: ScienceDirect/Algorithm_ws/l9-11.pdf\nOur RLS baseline: algorithms/core.py (parallel Gaussian local searches)',.72,3.55,11.85,1.5,21,MUTED)
    textbox(s,'Raw runs + config: results/workshop/  |  Method audit: docs/methodology.md\nRebuild: python -m workshop all  |  Thai script + Q&A: presentation/',.72,5.7,11.85,1.0,21,TEAL,True)
    # Speaker notes: timing and artifact pointers travel with the deck.
    for i,s in enumerate(prs.slides):
        s.notes_slide.notes_text_frame.text=f'Slide {i+1}: {TITLES[i]}\nAllocated time: {TIMES[i]} seconds.\nUse presentation/script_th.md, section {i+1}.\nAll performance claims use results/workshop/summary.json.\nSee docs/methodology.md for normalization and fair-budget details.'
    pptx=output/'GOA_Optimization_Workshop.pptx';prs.save(pptx)
    save_json(output/'slide_plan.json',dict(slides=[dict(slide=i+1,title=t,seconds=TIMES[i]) for i,t in enumerate(TITLES)],total_seconds=sum(TIMES),qa_included_in_talk_time=False))
    # Rebuild narrative files as part of the presentation command.
    from workshop.narrative import write_narrative
    write_narrative(root,output)
    conversion=dict(pptx=str(pptx),pdf_available=False)
    if shutil.which('libreoffice'):
        # Convert into a fresh directory so an older PDF cannot mask failure.
        with tempfile.TemporaryDirectory(prefix='goa-workshop-pdf-') as temporary:
            command=['libreoffice','-env:UserInstallation=file:///tmp/goa-workshop-lo','--headless','--convert-to','pdf','--outdir',temporary,str(pptx.resolve())]
            try:
                process=subprocess.run(command,capture_output=True,text=True,timeout=120)
                conversion.update(command=command,returncode=process.returncode,output=process.stdout+process.stderr)
                generated=Path(temporary)/'GOA_Optimization_Workshop.pdf'
                if process.returncode==0 and generated.exists():
                    shutil.copy2(generated,output/generated.name)
                    conversion['pdf_available']=True
                    conversion['pptx_sha256']=hashlib.sha256(pptx.read_bytes()).hexdigest()
                    conversion['pdf_sha256']=hashlib.sha256((output/generated.name).read_bytes()).hexdigest()
            except (OSError,subprocess.TimeoutExpired) as error:
                conversion['reason']=str(error)
    else: conversion['reason']='LibreOffice not installed; use PPTX or PNG slide exports after installing LibreOffice.'
    save_json(output/'export_status.json',conversion)
    if not conversion['pdf_available']:print('PDF conversion unavailable:',conversion,flush=True)
    else:print('12-slide PPTX + PDF generated',flush=True)
    return pptx
